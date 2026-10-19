import base64
import io
from collections.abc import Callable

import pytest
from docx import Document
from fastapi.testclient import TestClient
from happy_service import MAX_BYTES, app

TOKEN = "test-machine-token-with-at-least-thirty-two-characters"
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def pdf_fixture() -> bytes:
    stream = (
        b"BT /F1 12 Tf 72 720 Td (Ama Mensah) Tj "
        b"0 -18 Td (ama@example.org) Tj "
        b"0 -18 Td (0241234567) Tj "
        b"0 -18 Td (Customer service and Excel skills) Tj ET\n"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"endstream",
    ]
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, value in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{number} 0 obj\n".encode())
        content.extend(value + b"\nendobj\n")
    xref_offset = len(content)
    content.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        content.extend(f"{offset:010} 00000 n \n".encode())
    content.extend(
        b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n"
        + str(xref_offset).encode()
        + b"\n%%EOF\n"
    )
    return bytes(content)


def docx_fixture() -> bytes:
    document = Document()
    document.add_paragraph("Ama Mensah")
    document.add_paragraph("ama@example.org")
    document.add_paragraph("Customer service and Excel skills")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def payload(filename: str, mime_type: str, content: bytes) -> dict[str, str]:
    return {
        "filename": filename,
        "mime_type": mime_type,
        "content_base64": base64.b64encode(content).decode("ascii"),
    }


@pytest.fixture
def client(monkeypatch: pytest.MonkeyPatch) -> TestClient:
    monkeypatch.setenv("HAPPY_PARSER_TOKEN", TOKEN)
    return TestClient(app)


@pytest.mark.parametrize(
    ("filename", "mime_type", "content"),
    [
        ("candidate.pdf", PDF_MIME, pdf_fixture),
        ("candidate.docx", DOCX_MIME, docx_fixture),
    ],
)
def test_happy_parse_returns_advisory_suggestions_without_pii_logs(
    client: TestClient,
    caplog: pytest.LogCaptureFixture,
    filename: str,
    mime_type: str,
    content: Callable[[], bytes],
) -> None:
    response = client.post(
        "/api/v1/integrations/happy/parse",
        headers={"Authorization": f"Bearer {TOKEN}"},
        json=payload(filename, mime_type, content()),
    )

    assert response.status_code == 200, response.text
    assert response.headers["Cache-Control"] == "no-store"
    result = response.json()
    assert result["schema_version"] == "1"
    assert result["provider"] == "happy-cv-parser"
    assert result["suggestions"]
    assert any(item["field"] == "email" for item in result["suggestions"])
    assert "ama@example.org" not in caplog.text
    assert "Ama Mensah" not in caplog.text


def test_machine_authentication_is_required(client: TestClient) -> None:
    response = client.post(
        "/api/v1/integrations/happy/parse",
        json=payload("candidate.pdf", PDF_MIME, pdf_fixture()),
    )

    assert response.status_code == 401
    assert response.json() == {"detail": "Valid machine credentials are required."}


@pytest.mark.parametrize(
    "request_payload",
    [
        payload("candidate.pdf", PDF_MIME, b"not a PDF"),
        payload("candidate.pdf", DOCX_MIME, b"not a DOCX"),
        {
            "filename": "candidate.pdf",
            "mime_type": PDF_MIME,
            "content_base64": "not-base64!",
        },
        payload("candidate.exe", "application/octet-stream", b"bad"),
    ],
)
def test_malformed_or_unsupported_uploads_return_safe_validation_errors(
    client: TestClient, request_payload: dict[str, str]
) -> None:
    response = client.post(
        "/api/v1/integrations/happy/parse",
        headers={"Authorization": f"Bearer {TOKEN}"},
        json=request_payload,
    )

    assert response.status_code == 422
    assert "ama@example.org" not in response.text


def test_request_size_is_bounded_and_rejected_responses_are_not_cached(
    client: TestClient,
) -> None:
    response = client.post(
        "/api/v1/integrations/happy/parse",
        headers={"Authorization": f"Bearer {TOKEN}", "Content-Type": "application/json"},
        content=b" " * (14 * 1024 * 1024 + 1),
    )

    assert response.status_code == 413
    assert response.headers["Cache-Control"] == "no-store"
    assert MAX_BYTES == 10 * 1024 * 1024
